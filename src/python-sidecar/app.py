from __future__ import annotations

import hashlib
import math
import sqlite3
import threading
import time
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

import numpy as np
from bs4 import BeautifulSoup
from docx import Document
from fastapi import FastAPI, HTTPException
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer
from turbovec import IdMapIndex

APP_DIR = Path(__file__).resolve().parent
STORE_DIR = APP_DIR / "rag_store"
INDEX_PATH = STORE_DIR / "index.tvim"
DB_PATH = STORE_DIR / "metadata.db"
EMBEDDING_MODEL = "BAAI/bge-small-en-v1.5"
VECTOR_DIMENSION = 384
NORMALIZE_EMBEDDINGS = True
DISTANCE_METRIC = "cosine"
QUERY_INSTRUCTION_TEMPLATE = "Represent this sentence for searching relevant passages: {query}"
BIT_WIDTH = 4
SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx", ".html", ".htm"}

app = FastAPI(title="Local RAG POC Sidecar")
_store_lock = threading.RLock()
_model_lock = threading.Lock()
_model: SentenceTransformer | None = None
_index: IdMapIndex | None = None


class IngestRequest(BaseModel):
    folder_path: str = Field(alias="folderPath")
    chunk_size: int = Field(default=900, alias="chunkSize", gt=0)
    overlap: int = Field(default=150, ge=0)
    allowed_extensions: list[str] = Field(default_factory=list, alias="allowedExtensions")

    model_config = {"populate_by_name": True}


class IngestResponse(BaseModel):
    documents_discovered: int = Field(alias="documentsDiscovered")
    documents_indexed: int = Field(alias="documentsIndexed")
    documents_skipped: int = Field(alias="documentsSkipped")
    chunks_indexed: int = Field(alias="chunksIndexed")
    elapsed_seconds: float = Field(alias="elapsedSeconds")
    index_path: str = Field(alias="indexPath")
    metadata_db_path: str = Field(alias="metadataDbPath")
    message: str | None = None

    model_config = {"populate_by_name": True}


class SearchRequest(BaseModel):
    query: str
    top_k: int = Field(default=10, alias="topK", gt=0)

    model_config = {"populate_by_name": True}


class EmbedRequest(BaseModel):
    text: str
    apply_query_prefix: bool = Field(default=True, alias="applyQueryPrefix")

    model_config = {"populate_by_name": True}


class EmbedResponse(BaseModel):
    embedding: list[float]
    embedding_model: str = Field(alias="embeddingModel")
    vector_dimension: int = Field(alias="vectorDimension")
    normalize_embeddings: bool = Field(alias="normalizeEmbeddings")
    distance_metric: str = Field(alias="distanceMetric")

    model_config = {"populate_by_name": True}


class SearchResult(BaseModel):
    score: float
    document_path: str = Field(alias="documentPath")
    chunk_index: int = Field(alias="chunkIndex")
    text: str

    model_config = {"populate_by_name": True}


class SearchResponse(BaseModel):
    results: list[SearchResult]
    elapsed_milliseconds: float = Field(alias="elapsedMilliseconds")
    chunks_searched: int = Field(alias="chunksSearched")

    model_config = {"populate_by_name": True}


class BenchmarkRequest(BaseModel):
    query: str
    top_k: int = Field(default=10, alias="topK", gt=0)
    runs: int = Field(default=25, gt=0, le=500)

    model_config = {"populate_by_name": True}


class BenchmarkResponse(BaseModel):
    average_milliseconds: float = Field(alias="averageMilliseconds")
    p95_milliseconds: float = Field(alias="p95Milliseconds")
    runs: int
    chunks_searched: int = Field(alias="chunksSearched")

    model_config = {"populate_by_name": True}


class StatsResponse(BaseModel):
    document_count: int = Field(alias="documentCount")
    chunk_count: int = Field(alias="chunkCount")
    index_path: str = Field(alias="indexPath")
    metadata_db_path: str = Field(alias="metadataDbPath")
    index_exists: bool = Field(alias="indexExists")
    embedding_model: str = Field(alias="embeddingModel")
    vector_dimension: int = Field(alias="vectorDimension")
    bit_width: int = Field(alias="bitWidth")
    normalize_embeddings: bool = Field(alias="normalizeEmbeddings")
    distance_metric: str = Field(alias="distanceMetric")

    model_config = {"populate_by_name": True}


def get_connection() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db() -> None:
    STORE_DIR.mkdir(parents=True, exist_ok=True)
    with get_connection() as conn:
        conn.executescript(
            """
            CREATE TABLE IF NOT EXISTS documents (
                id INTEGER PRIMARY KEY,
                path TEXT NOT NULL UNIQUE,
                filename TEXT NOT NULL,
                extension TEXT NOT NULL,
                sha256 TEXT NOT NULL,
                modified_utc TEXT NOT NULL,
                indexed_utc TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS chunks (
                id INTEGER PRIMARY KEY,
                document_id INTEGER NOT NULL,
                chunk_index INTEGER NOT NULL,
                vector_id INTEGER NOT NULL UNIQUE,
                text TEXT NOT NULL,
                char_start INTEGER,
                char_end INTEGER,
                FOREIGN KEY(document_id) REFERENCES documents(id) ON DELETE CASCADE
            );

            CREATE TABLE IF NOT EXISTS embedding_metadata (
                id INTEGER PRIMARY KEY CHECK (id = 1),
                embedding_model_id TEXT NOT NULL,
                embedding_dimensions INTEGER NOT NULL,
                normalize_embeddings INTEGER NOT NULL CHECK (normalize_embeddings IN (0, 1)),
                distance_metric TEXT NOT NULL
            );
            """
        )
        validate_embedding_metadata(conn)


def validate_embedding_metadata(conn: sqlite3.Connection) -> None:
    row = conn.execute(
        """
        SELECT embedding_model_id, embedding_dimensions, normalize_embeddings, distance_metric
        FROM embedding_metadata
        WHERE id = 1
        """
    ).fetchone()
    expected = (EMBEDDING_MODEL, VECTOR_DIMENSION, int(NORMALIZE_EMBEDDINGS), DISTANCE_METRIC)
    if row is None:
        conn.execute(
            """
            INSERT INTO embedding_metadata(
                id, embedding_model_id, embedding_dimensions, normalize_embeddings, distance_metric
            )
            VALUES (1, ?, ?, ?, ?)
            """,
            expected,
        )
        return

    actual = (row["embedding_model_id"], row["embedding_dimensions"], row["normalize_embeddings"], row["distance_metric"])
    if actual != expected:
        raise RuntimeError(
            "Embedding metadata mismatch: expected "
            f"model={EMBEDDING_MODEL}, dimensions={VECTOR_DIMENSION}, "
            f"normalize_embeddings={NORMALIZE_EMBEDDINGS}, distance_metric={DISTANCE_METRIC}; "
            f"found model={actual[0]}, dimensions={actual[1]}, "
            f"normalize_embeddings={bool(actual[2])}, distance_metric={actual[3]}. "
            "Rebuild/reingest the vector store with the configured embedding model."
        )


def format_query_for_embedding(query: str) -> str:
    return QUERY_INSTRUCTION_TEMPLATE.format(query=query)


def load_or_create_index() -> IdMapIndex:
    STORE_DIR.mkdir(parents=True, exist_ok=True)
    if INDEX_PATH.exists():
        return IdMapIndex.load(str(INDEX_PATH))
    return IdMapIndex(dim=VECTOR_DIMENSION, bit_width=BIT_WIDTH)


def get_index() -> IdMapIndex:
    global _index
    if _index is None:
        _index = load_or_create_index()
    return _index


def get_model() -> SentenceTransformer:
    global _model
    if _model is None:
        with _model_lock:
            if _model is None:
                _model = SentenceTransformer(EMBEDDING_MODEL)
    return _model


def normalize_extensions(extensions: Iterable[str]) -> set[str]:
    normalized = set()
    for extension in extensions:
        value = extension.strip().lower()
        if not value:
            continue
        if not value.startswith("."):
            value = "." + value
        if value in SUPPORTED_EXTENSIONS:
            normalized.add(value)
    return normalized


def discover_files(folder: Path, extensions: set[str]) -> list[Path]:
    return sorted(path for path in folder.rglob("*") if path.is_file() and path.suffix.lower() in extensions)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def modified_utc(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()


def extract_text(path: Path) -> str:
    extension = path.suffix.lower()
    if extension == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if extension in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if extension == ".docx":
        document = Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs)
    if extension in {".html", ".htm"}:
        soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "lxml")
        return soup.get_text("\n")
    raise ValueError(f"Unsupported extension: {extension}")


def chunk_text(text: str, chunk_size: int, overlap: int) -> list[tuple[int, int, str]]:
    compact = text.replace("\r\n", "\n").strip()
    if not compact:
        return []
    chunks: list[tuple[int, int, str]] = []
    start = 0
    step = chunk_size - overlap
    while start < len(compact):
        end = min(start + chunk_size, len(compact))
        chunk = compact[start:end].strip()
        if chunk:
            chunks.append((start, end, chunk))
        if end >= len(compact):
            break
        start += step
    return chunks


def stable_vector_id(path: Path, modified: str, chunk_index: int, chunk: str) -> np.uint64:
    material = f"{path.resolve()}|{modified}|{chunk_index}|{hashlib.sha256(chunk.encode('utf-8')).hexdigest()}"
    digest = hashlib.sha256(material.encode("utf-8")).digest()
    return np.uint64(int.from_bytes(digest[:8], "big", signed=False) & 0x7FFFFFFFFFFFFFFF)


def embed_texts(texts: list[str]) -> np.ndarray:
    if not texts:
        return np.empty((0, VECTOR_DIMENSION), dtype=np.float32)
    vectors = get_model().encode(texts, normalize_embeddings=NORMALIZE_EMBEDDINGS, convert_to_numpy=True, show_progress_bar=False)
    return np.asarray(vectors, dtype=np.float32)


def add_vectors(index: IdMapIndex, vectors: np.ndarray, ids: list[np.uint64]) -> None:
    if len(ids) == 0:
        return
    index.add_with_ids(np.asarray(vectors, dtype=np.float32), np.asarray(ids, dtype=np.uint64))


def remove_existing_document(conn: sqlite3.Connection, index: IdMapIndex, document_id: int) -> None:
    rows = conn.execute("SELECT vector_id FROM chunks WHERE document_id = ?", (document_id,)).fetchall()
    for row in rows:
        try:
            index.remove(np.uint64(row["vector_id"]))
        except Exception:
            pass
    conn.execute("DELETE FROM chunks WHERE document_id = ?", (document_id,))
    conn.execute("DELETE FROM documents WHERE id = ?", (document_id,))


def perform_search(query: str, top_k: int) -> SearchResponse:
    if not query.strip():
        raise HTTPException(status_code=400, detail="Query is required.")

    with _store_lock:
        index = get_index()
        chunk_count = len(index)
        if chunk_count == 0:
            return SearchResponse(results=[], elapsedMilliseconds=0.0, chunksSearched=0)
        query_vector = embed_texts([format_query_for_embedding(query)])
        k = min(top_k, chunk_count)
        start = time.perf_counter()
        scores, ids = index.search(query_vector, k)
        elapsed_ms = (time.perf_counter() - start) * 1000.0
        vector_ids = [int(value) for value in ids[0].tolist()]
        score_values = [float(value) for value in scores[0].tolist()]

    if not vector_ids:
        return SearchResponse(results=[], elapsedMilliseconds=elapsed_ms, chunksSearched=chunk_count)

    placeholders = ",".join("?" for _ in vector_ids)
    with get_connection() as conn:
        rows = conn.execute(
            f"""
            SELECT c.vector_id, c.chunk_index, c.text, d.path
            FROM chunks c
            JOIN documents d ON d.id = c.document_id
            WHERE c.vector_id IN ({placeholders})
            """,
            vector_ids,
        ).fetchall()
    by_id = {int(row["vector_id"]): row for row in rows}
    results = []
    for vector_id, score in zip(vector_ids, score_values):
        row = by_id.get(vector_id)
        if row is None:
            continue
        results.append(
            SearchResult(score=score, documentPath=row["path"], chunkIndex=row["chunk_index"], text=row["text"])
        )
    return SearchResponse(results=results, elapsedMilliseconds=elapsed_ms, chunksSearched=chunk_count)


@app.on_event("startup")
def startup() -> None:
    init_db()
    get_index()


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "OK"}


@app.get("/stats", response_model=StatsResponse, response_model_by_alias=True)
def stats() -> StatsResponse:
    init_db()
    with get_connection() as conn:
        document_count = conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0]
        chunk_count = conn.execute("SELECT COUNT(*) FROM chunks").fetchone()[0]
    return StatsResponse(
        documentCount=document_count,
        chunkCount=chunk_count,
        indexPath=str(INDEX_PATH),
        metadataDbPath=str(DB_PATH),
        indexExists=INDEX_PATH.exists(),
        embeddingModel=EMBEDDING_MODEL,
        vectorDimension=VECTOR_DIMENSION,
        bitWidth=BIT_WIDTH,
        normalizeEmbeddings=NORMALIZE_EMBEDDINGS,
        distanceMetric=DISTANCE_METRIC,
    )


@app.post("/ingest-folder", response_model=IngestResponse, response_model_by_alias=True)
def ingest_folder(request: IngestRequest) -> IngestResponse:
    if request.overlap >= request.chunk_size:
        raise HTTPException(status_code=400, detail="Overlap must be smaller than chunk size.")
    folder = Path(request.folder_path).expanduser().resolve()
    if not folder.exists() or not folder.is_dir():
        raise HTTPException(status_code=400, detail=f"Folder does not exist: {folder}")
    extensions = normalize_extensions(request.allowed_extensions) or SUPPORTED_EXTENSIONS
    files = discover_files(folder, extensions)
    started = time.perf_counter()
    indexed_documents = 0
    skipped_documents = 0
    indexed_chunks = 0

    with _store_lock:
        index = get_index()
        with get_connection() as conn:
            for path in files:
                digest = sha256_file(path)
                existing = conn.execute("SELECT id, sha256 FROM documents WHERE path = ?", (str(path),)).fetchone()
                if existing and existing["sha256"] == digest:
                    skipped_documents += 1
                    continue
                if existing:
                    remove_existing_document(conn, index, int(existing["id"]))

                text = extract_text(path)
                chunks = chunk_text(text, request.chunk_size, request.overlap)
                if not chunks:
                    skipped_documents += 1
                    continue

                modified = modified_utc(path)
                now = datetime.now(timezone.utc).isoformat()
                cursor = conn.execute(
                    """
                    INSERT INTO documents(path, filename, extension, sha256, modified_utc, indexed_utc)
                    VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (str(path), path.name, path.suffix.lower(), digest, modified, now),
                )
                document_id = int(cursor.lastrowid)
                chunk_texts = [chunk for _, _, chunk in chunks]
                vector_ids = [stable_vector_id(path, modified, i, chunk) for i, (_, _, chunk) in enumerate(chunks)]
                vectors = embed_texts(chunk_texts)
                add_vectors(index, vectors, vector_ids)
                for chunk_index, ((char_start, char_end, chunk), vector_id) in enumerate(zip(chunks, vector_ids)):
                    conn.execute(
                        """
                        INSERT INTO chunks(document_id, chunk_index, vector_id, text, char_start, char_end)
                        VALUES (?, ?, ?, ?, ?, ?)
                        """,
                        (document_id, chunk_index, int(vector_id), chunk, char_start, char_end),
                    )
                indexed_documents += 1
                indexed_chunks += len(chunks)
            conn.commit()
        index.write(str(INDEX_PATH))

    elapsed = time.perf_counter() - started
    return IngestResponse(
        documentsDiscovered=len(files),
        documentsIndexed=indexed_documents,
        documentsSkipped=skipped_documents,
        chunksIndexed=indexed_chunks,
        elapsedSeconds=elapsed,
        indexPath=str(INDEX_PATH),
        metadataDbPath=str(DB_PATH),
        message="Ingest complete.",
    )


@app.post("/embed", response_model=EmbedResponse, response_model_by_alias=True)
def embed(request: EmbedRequest) -> EmbedResponse:
    if not request.text.strip():
        raise HTTPException(status_code=400, detail="Text is required.")
    input_text = format_query_for_embedding(request.text) if request.apply_query_prefix else request.text
    vector = embed_texts([input_text])[0].astype(float).tolist()
    return EmbedResponse(
        embedding=vector,
        embeddingModel=EMBEDDING_MODEL,
        vectorDimension=VECTOR_DIMENSION,
        normalizeEmbeddings=NORMALIZE_EMBEDDINGS,
        distanceMetric=DISTANCE_METRIC,
    )


@app.post("/search", response_model=SearchResponse, response_model_by_alias=True)
def search(request: SearchRequest) -> SearchResponse:
    return perform_search(request.query, request.top_k)


@app.post("/benchmark", response_model=BenchmarkResponse, response_model_by_alias=True)
def benchmark(request: BenchmarkRequest) -> BenchmarkResponse:
    perform_search(request.query, request.top_k)  # warmup: loads model and primes search path
    timings = []
    chunks_searched = 0
    for _ in range(request.runs):
        response = perform_search(request.query, request.top_k)
        timings.append(response.elapsed_milliseconds)
        chunks_searched = response.chunks_searched
    timings.sort()
    average = sum(timings) / len(timings)
    p95_index = min(len(timings) - 1, math.ceil(len(timings) * 0.95) - 1)
    return BenchmarkResponse(
        averageMilliseconds=average,
        p95Milliseconds=timings[p95_index],
        runs=request.runs,
        chunksSearched=chunks_searched,
    )
